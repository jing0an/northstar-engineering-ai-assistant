from pathlib import Path


class DocumentParseError(ValueError):
    """Raised when a supported document cannot be parsed or contains no text."""


def _clean_surrogates(text: str) -> str:
    """Remove lone Unicode surrogate code points from extracted text."""
    return "".join(c for c in text if not 0xD800 <= ord(c) <= 0xDFFF)


def extract_pages(
    file_path: str | Path, file_type: str
) -> list[dict[str, object]]:
    """Extract text while preserving PDF page boundaries.

    DOCX content is returned as one logical content unit because python-docx
    does not expose reliable rendered physical page boundaries.
    """
    path = Path(file_path)
    normalized_type = file_type.lower().strip().lstrip(".")
    if normalized_type not in {"pdf", "docx"}:
        raise DocumentParseError(
            f"Unsupported document type: {file_type}. Only pdf and docx are supported."
        )
    if not path.exists() or not path.is_file():
        raise DocumentParseError(f"Document file does not exist: {path}")

    try:
        if normalized_type == "pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            pages = [
                {
                    "page_number": page_number,
                    "text": _clean_surrogates(page.extract_text() or ""),
                }
                for page_number, page in enumerate(reader.pages, start=1)
            ]
        else:
            from docx import Document

            document = Document(str(path))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            table_cells = [
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ]
            pages = [
                {
                    "page_number": 1,
                    "text": _clean_surrogates(
                        "\n".join(paragraphs + table_cells)
                    ),
                }
            ]
    except Exception as error:
        raise DocumentParseError(f"Failed to parse document: {path}") from error

    if not "\n".join(str(page["text"]) for page in pages).strip():
        raise DocumentParseError(f"Document contains no extractable text: {path}")
    return pages


def extract_text(file_path: str | Path, file_type: str) -> str:
    """Extract text from a PDF or DOCX file using the matching parser."""
    pages = extract_pages(file_path, file_type)
    return "\n".join(str(page["text"]) for page in pages).strip()