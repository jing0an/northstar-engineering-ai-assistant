"""Concrete parsers for common text and office document formats."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from .base import DocumentParser


def _page(
    page_number: int,
    text: str,
    *,
    metadata: dict[str, Any] | None = None,
    blocks: list[dict[str, Any]] | None = None,
) -> dict[str, object]:
    result: dict[str, object] = {"page_number": page_number, "text": text}
    if metadata is not None:
        result["metadata"] = metadata
    if blocks is not None:
        result["blocks"] = blocks
    return result


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig")


class TxtParser(DocumentParser):
    supported_extensions = frozenset({".txt"})

    def parse(self, file_path: str | Path, file_type: str | None = None) -> list[dict[str, object]]:
        path = Path(file_path)
        text = _read_text(path)
        return [_page(1, text, metadata={"format": "txt"})]


class MdParser(DocumentParser):
    supported_extensions = frozenset({".md"})

    def parse(self, file_path: str | Path, file_type: str | None = None) -> list[dict[str, object]]:
        path = Path(file_path)
        text = _read_text(path)
        blocks: list[dict[str, Any]] = []
        lines = text.replace("\r\n", "\n").replace("\r", "\n").splitlines()
        in_code = False
        code_lines: list[str] = []
        for line in lines:
            stripped = line.strip()
            if stripped.startswith("```") or stripped.startswith("~~~"):
                if in_code:
                    code_lines.append(line)
                    blocks.append({"type": "code", "text": "\n".join(code_lines)})
                    code_lines = []
                    in_code = False
                else:
                    in_code = True
                    code_lines = [line]
                continue
            if in_code:
                code_lines.append(line)
                continue
            if stripped.startswith("#"):
                blocks.append({"type": "heading", "level": len(stripped) - len(stripped.lstrip("#")), "text": line})
            elif stripped.startswith(">"):
                blocks.append({"type": "quote", "text": line})
            elif stripped.startswith(("- ", "* ", "+ ")) or (len(stripped) > 2 and stripped[0].isdigit() and stripped[1:3] == ". "):
                blocks.append({"type": "list_item", "text": line})
            elif stripped:
                blocks.append({"type": "paragraph", "text": line})
        if in_code and code_lines:
            blocks.append({"type": "code", "text": "\n".join(code_lines)})
        if "|" in text and any(line.count("|") >= 2 for line in lines):
            blocks.append({"type": "table", "text": "\n".join(line for line in lines if "|" in line)})
        return [_page(1, text, metadata={"format": "md"}, blocks=blocks)]


class DocxParser(DocumentParser):
    supported_extensions = frozenset({".docx"})

    def parse(self, file_path: str | Path, file_type: str | None = None) -> list[dict[str, object]]:
        from docx import Document
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        from docx.table import Table, _Cell
        from docx.text.paragraph import Paragraph

        document = Document(str(file_path))
        blocks: list[dict[str, Any]] = []
        text_parts: list[str] = []
        parent = document.element.body
        for element in parent.iterchildren():
            if isinstance(element, CT_P):
                paragraph = Paragraph(element, document)
                text = paragraph.text
                style_name = paragraph.style.name if paragraph.style is not None else ""
                style_lower = style_name.casefold()
                block_type = "heading" if style_lower.startswith("heading") else "paragraph"
                if "list" in style_lower:
                    block_type = "list_item"
                level = None
                if block_type == "heading":
                    try:
                        level = int(style_name.split()[-1])
                    except (ValueError, IndexError):
                        level = None
                blocks.append({"type": block_type, "text": text, "style": style_name, "level": level})
                if text:
                    text_parts.append(text)
            elif isinstance(element, CT_Tbl):
                table = Table(element, document)
                rows: list[list[str]] = []
                for row in table.rows:
                    cells = [cell.text for cell in row.cells]
                    rows.append(cells)
                    text_parts.append(" | ".join(cells))
                blocks.append({"type": "table", "rows": rows})
        image_count = len(document.inline_shapes)
        metadata: dict[str, Any] = {"format": "docx", "image_count": image_count}
        if image_count:
            blocks.append({"type": "image", "count": image_count})
        return [_page(1, "\n".join(text_parts), metadata=metadata, blocks=blocks)]


class PptxParser(DocumentParser):
    supported_extensions = frozenset({".pptx"})

    def parse(self, file_path: str | Path, file_type: str | None = None) -> list[dict[str, object]]:
        from pptx import Presentation

        presentation = Presentation(str(file_path))
        pages: list[dict[str, object]] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            blocks: list[dict[str, Any]] = []
            text_parts: list[str] = []
            title: str | None = None
            image_count = 0
            for shape in slide.shapes:
                try:
                    if getattr(shape, "has_table", False):
                        rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                        blocks.append({"type": "table", "rows": rows})
                        text_parts.extend(" | ".join(row) for row in rows)
                    elif getattr(shape, "shape_type", None) == 13:
                        image_count += 1
                        blocks.append({"type": "image"})
                    elif getattr(shape, "has_text_frame", False):
                        text = shape.text
                        if not text:
                            continue
                        is_title = bool(getattr(shape, "is_placeholder", False) and getattr(shape.placeholder_format, "type", None) in (1, 3))
                        if is_title and title is None:
                            title = text
                        blocks.append({"type": "title" if is_title else "text", "text": text})
                        text_parts.append(text)
                except Exception as error:
                    blocks.append({"type": "unreadable", "error": type(error).__name__})
            metadata = {"format": "pptx", "slide_number": slide_number, "title": title, "image_count": image_count}
            pages.append(_page(slide_number, "\n".join(text_parts), metadata=metadata, blocks=blocks))
        return pages


class XlsxParser(DocumentParser):
    supported_extensions = frozenset({".xlsx"})

    def parse(self, file_path: str | Path, file_type: str | None = None) -> list[dict[str, object]]:
        from openpyxl import load_workbook

        workbook = load_workbook(str(file_path), data_only=False, read_only=False)
        pages: list[dict[str, object]] = []
        for page_number, worksheet in enumerate(workbook.worksheets, start=1):
            rows: list[dict[str, Any]] = []
            text_rows: list[str] = []
            for row in worksheet.iter_rows():
                cells = []
                for cell in row:
                    value = cell.value
                    if value is None:
                        text_value = ""
                    elif hasattr(value, "isoformat") and not isinstance(value, str):
                        text_value = value.isoformat()
                    else:
                        text_value = str(value)
                    cells.append({"row": cell.row, "column": cell.column, "coordinate": cell.coordinate, "value": value, "text": text_value})
                if any(cell["value"] is not None for cell in cells):
                    rows.append({"row": row[0].row if row else 0, "cells": cells})
                    text_rows.append("\t".join(cell["text"] for cell in cells).rstrip("\t"))
            blocks = [{"type": "row", **row} for row in rows]
            pages.append(_page(page_number, "\n".join(text_rows), metadata={"format": "xlsx", "sheet_name": worksheet.title}, blocks=blocks))
        return pages