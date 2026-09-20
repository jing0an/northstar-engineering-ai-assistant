import io
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

import app.api.routes.files as files_route
from app.auth.models import ProjectMembership, ProjectRole, User, UserStatus
from app.auth.persistence import SQLiteProjectMembershipRepository, SQLiteUserRepository
from app.auth.store import AuthStore
from app.auth.token import AccessTokenService
from app.main import app
from app.services.file_storage import FileStorage


class FilePreviewTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.storage_root = Path(self.temp_dir.name) / "storage"
        self.auth_db = Path(self.temp_dir.name) / "auth.sqlite3"
        os.environ["ACCESS_TOKEN_SECRET"] = "test-secret-for-file-preview-1234567890"

        self.auth_store = AuthStore(self.auth_db)
        self.user_repository = SQLiteUserRepository(self.auth_store)
        self.membership_repository = SQLiteProjectMembershipRepository(self.auth_store)
        self.user = User(
            username="file-preview-user",
            email="file-preview@example.com",
            password_hash="unused",
            status=UserStatus.ACTIVE,
            email_verified=True,
        )
        self.user_repository.create(self.user)
        for project_id in ("project-a", "project-b"):
            self.membership_repository.add(
                ProjectMembership(
                    user_id=self.user.user_id,
                    project_id=project_id,
                    role=ProjectRole.MEMBER,
                )
            )
        self.access_token = AccessTokenService().create_access_token(self.user.user_id)
        self.original_storage = files_route.file_storage
        files_route.file_storage = FileStorage(self.storage_root)
        self.auth_store_patcher = patch(
            "app.auth.dependencies.AuthStore", lambda: AuthStore(self.auth_db)
        )
        self.auth_store_patcher.start()
        self.client = TestClient(app)

    def tearDown(self) -> None:
        self.auth_store_patcher.stop()
        files_route.file_storage = self.original_storage
        self.auth_store.close()
        self.temp_dir.cleanup()

    @property
    def headers(self) -> dict[str, str]:
        return {"Authorization": f"Bearer {self.access_token}"}

    def add_file(self, project_id: str, *, file_id: str, file_type: str, content: bytes) -> None:
        project_dir = self.storage_root / project_id
        project_dir.mkdir(parents=True, exist_ok=True)
        stored_filename = f"{file_id}.{file_type}"
        (project_dir / stored_filename).write_bytes(content)
        (project_dir / "metadata.json").write_text(
            json.dumps([{
                "file_id": file_id,
                "project_id": project_id,
                "original_filename": f"{file_id}.{file_type}",
                "stored_filename": stored_filename,
                "file_type": file_type,
                "document_version": 4,
                "is_current": True,
            }], ensure_ascii=False),
            encoding="utf-8",
        )

    @staticmethod
    def make_docx() -> bytes:
        document = Document()
        document.add_heading("项目概况", level=1)
        document.add_paragraph("这是正文段落。")
        document.add_paragraph("第一项", style="List Bullet")
        document.add_paragraph("第二项", style="List Bullet")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "字段"
        table.cell(0, 1).text = "值"
        table.cell(1, 0).text = "状态"
        table.cell(1, 1).text = "正常"
        stream = io.BytesIO()
        document.save(stream)
        return stream.getvalue()

    def preview(self, project_id: str, file_id: str = "brief"):
        return self.client.get(
            f"/api/files/{file_id}/preview",
            params={"project_id": project_id},
            headers=self.headers,
        )

    def test_authenticated_owner_gets_docx_preview_json(self) -> None:
        self.add_file("project-a", file_id="brief", file_type="docx", content=self.make_docx())
        response = self.preview("project-a")

        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["file_id"], "brief")
        self.assertEqual(body["project_id"], "project-a")
        self.assertEqual(body["document_version"], 4)
        self.assertEqual(body["file_type"], "docx")
        self.assertIsInstance(body["blocks"], list)
        self.assertEqual(body["blocks"][0], {"type": "heading", "level": 1, "text": "项目概况"})
        self.assertIn({"type": "paragraph", "text": "这是正文段落。"}, body["blocks"])
        self.assertIn({"type": "list", "items": ["第一项", "第二项"]}, body["blocks"])
        self.assertIn({"type": "table", "rows": [["字段", "值"], ["状态", "正常"]]}, body["blocks"])

    def test_preview_requires_authentication(self) -> None:
        self.add_file("project-a", file_id="brief", file_type="docx", content=self.make_docx())
        response = self.client.get("/api/files/brief/preview", params={"project_id": "project-a"})
        self.assertEqual(response.status_code, 401)

    def test_preview_requires_project_membership(self) -> None:
        self.add_file("project-c", file_id="brief", file_type="docx", content=self.make_docx())
        self.assertEqual(self.preview("project-c").status_code, 403)

    def test_file_must_belong_to_requested_project(self) -> None:
        self.add_file("project-b", file_id="other", file_type="docx", content=self.make_docx())
        self.assertEqual(self.preview("project-a", "other").status_code, 404)

    def test_missing_file_returns_404(self) -> None:
        self.assertEqual(self.preview("project-a").status_code, 404)

    def test_pdf_preview_is_explicitly_unsupported(self) -> None:
        self.add_file("project-a", file_id="brief", file_type="pdf", content=b"%PDF-1.4")
        response = self.preview("project-a")
        self.assertEqual(response.status_code, 415)

    def test_metadata_path_traversal_is_rejected(self) -> None:
        project_dir = self.storage_root / "project-a"
        project_dir.mkdir(parents=True, exist_ok=True)
        (self.storage_root / "outside.docx").write_bytes(self.make_docx())
        (project_dir / "metadata.json").write_text(json.dumps([{
            "file_id": "brief",
            "project_id": "project-a",
            "original_filename": "brief.docx",
            "stored_filename": "../outside.docx",
            "file_type": "docx",
            "document_version": 1,
        }]), encoding="utf-8")
        self.assertEqual(self.preview("project-a").status_code, 404)


if __name__ == "__main__":
    unittest.main()
