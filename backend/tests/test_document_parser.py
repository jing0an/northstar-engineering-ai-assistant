import tempfile
import unittest
from pathlib import Path

from docx import Document

from app.services.document_parser import (
    DocumentParseError,
    extract_pages,
    extract_text,
)


def build_test_pdf(text: str) -> bytes:
    return build_multi_page_pdf([text])


def build_multi_page_pdf(texts: list[str]) -> bytes:
    page_numbers = [3 + index * 2 for index in range(len(texts))]
    content_numbers = [page_number + 1 for page_number in page_numbers]
    font_number = 3 + len(texts) * 2
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        (
            b"<< /Type /Pages /Kids ["
            + b" ".join(
                f"{page_number} 0 R".encode("ascii")
                for page_number in page_numbers
            )
            + f"] /Count {len(texts)} >>".encode("ascii")
        ),
    ]

    for content_number, text in zip(content_numbers, texts):
        stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
        objects.extend(
            [
                (
                    b"<< /Type /Page /Parent 2 0 R "
                    b"/MediaBox [0 0 612 792] "
                    b"/Resources << /Font << /F1 "
                    + f"{font_number} 0 R".encode("ascii")
                    + b" >> >> /Contents "
                    + f"{content_number} 0 R >>".encode("ascii")
                ),
                (
                    b"<< /Length "
                    + str(len(stream)).encode("ascii")
                    + b" >>\nstream\n"
                    + stream
                    + b"\nendstream"
                ),
            ]
        )

    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode("ascii"))
        output.extend(obj)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    return bytes(output)


class DocumentParserTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_pdf_text_is_extracted(self) -> None:
        path = self.root / "sample.pdf"
        path.write_bytes(build_test_pdf("Northstar PDF test"))
        self.assertIn("Northstar PDF test", extract_text(path, "pdf"))

    def test_pdf_pages_preserve_page_numbers_and_text(self) -> None:
        path = self.root / "multi-page.pdf"
        path.write_bytes(build_multi_page_pdf(["Page one", "Page two"]))

        pages = extract_pages(path, "pdf")

        self.assertEqual([page["page_number"] for page in pages], [1, 2])
        self.assertIn("Page one", pages[0]["text"])
        self.assertIn("Page two", pages[1]["text"])

    def test_docx_text_is_extracted(self) -> None:
        path = self.root / "sample.docx"
        document = Document()
        document.add_paragraph("Northstar DOCX test")
        document.save(path)
        self.assertIn("Northstar DOCX test", extract_text(path, "docx"))

    def test_docx_pages_are_a_logical_content_unit(self) -> None:
        path = self.root / "sample.docx"
        document = Document()
        document.add_paragraph("Northstar DOCX test")
        document.save(path)

        pages = extract_pages(path, "docx")

        self.assertEqual(pages[0]["page_number"], 1)
        self.assertIn("Northstar DOCX test", pages[0]["text"])

    def test_unsupported_type_raises_clear_error(self) -> None:
        with self.assertRaisesRegex(DocumentParseError, "Unsupported document type"):
            extract_text(self.root / "sample.txt", "txt")

    def test_empty_text_raises_clear_error(self) -> None:
        path = self.root / "empty.docx"
        Document().save(path)
        with self.assertRaisesRegex(DocumentParseError, "no extractable text"):
            extract_text(path, "docx")


if __name__ == "__main__":
    unittest.main()