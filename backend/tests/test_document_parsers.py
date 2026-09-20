import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from app.services.document_parsers import (
    DocxParser,
    MdParser,
    PptxParser,
    TxtParser,
    XlsxParser,
    create_default_registry,
)


class DocumentParsersTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_txt_utf8_bom_and_empty_file(self) -> None:
        path = self.root / "notes.txt"
        path.write_bytes("中文文本\nsecond line".encode("utf-8-sig"))
        result = TxtParser().parse(path)
        self.assertEqual(result[0]["text"], "中文文本\nsecond line")
        empty = self.root / "empty.txt"
        empty.write_bytes(b"")
        self.assertEqual(TxtParser().parse(empty)[0]["text"], "")

    def test_markdown_preserves_source_and_classifies_blocks(self) -> None:
        path = self.root / "README.md"
        source = "# Heading\n\nParagraph with [link](https://example.com).\n- item\n\n> quote\n\n```python\nprint(1)\n```\n\n| A | B |\n|---|---|\n| 1 | 2 |\n"
        path.write_text(source, encoding="utf-8")
        result = MdParser().parse(path)[0]
        self.assertEqual(result["text"], source)
        block_types = [block["type"] for block in result["blocks"]]
        self.assertIn("heading", block_types)
        self.assertIn("paragraph", block_types)
        self.assertIn("list_item", block_types)
        self.assertIn("quote", block_types)
        self.assertIn("code", block_types)
        self.assertIn("table", block_types)

    def test_docx_preserves_order_headings_lists_and_table(self) -> None:
        path = self.root / "report.docx"
        document = Document()
        document.add_heading("Report", level=1)
        document.add_paragraph("First paragraph")
        document.add_paragraph("A list item", style="List Bullet")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Header"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "A"
        table.cell(1, 1).text = "1"
        document.add_paragraph("After table")
        document.save(path)
        page = DocxParser().parse(path)[0]
        self.assertIn("First paragraph", page["text"])
        self.assertIn("Header | Value", page["text"])
        self.assertEqual([block["type"] for block in page["blocks"]], ["heading", "paragraph", "list_item", "table", "paragraph"])
        self.assertEqual(page["blocks"][3]["rows"][1], ["A", "1"])

    def test_pptx_returns_one_page_per_slide_with_title_text_and_table(self) -> None:
        path = self.root / "deck.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Slide one"
        slide.placeholders[1].text = "Body text"
        table_shape = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1))
        table_shape.table.cell(0, 0).text = "A"
        table_shape.table.cell(0, 1).text = "B"
        table_shape.table.cell(1, 0).text = "1"
        table_shape.table.cell(1, 1).text = "2"
        presentation.slides.add_slide(presentation.slide_layouts[6])
        presentation.save(path)
        pages = PptxParser().parse(path)
        self.assertEqual(len(pages), 2)
        self.assertEqual(pages[0]["metadata"]["slide_number"], 1)
        self.assertEqual(pages[0]["metadata"]["title"], "Slide one")
        self.assertIn("Body text", pages[0]["text"])
        table_blocks = [block for block in pages[0]["blocks"] if block["type"] == "table"]
        self.assertEqual(table_blocks[0]["rows"][1], ["1", "2"])
        self.assertEqual(pages[1]["text"], "")

    def test_xlsx_returns_all_sheets_and_keeps_cells_and_formulas(self) -> None:
        path = self.root / "book.xlsx"
        workbook = Workbook()
        first = workbook.active
        first.title = "Summary"
        first.append(["Name", "Amount"])
        first.append(["A", 12])
        first["C1"] = "=SUM(B2)"
        workbook.create_sheet("Empty")
        second = workbook.create_sheet("Details")
        second.append(["Code", "Value"])
        second.append(["X", 3.5])
        workbook.save(path)
        pages = XlsxParser().parse(path)
        self.assertEqual([page["metadata"]["sheet_name"] for page in pages], ["Summary", "Empty", "Details"])
        summary = pages[0]
        self.assertIn("Name\tAmount", summary["text"])
        formula = next(cell for row in summary["blocks"] for cell in row["cells"] if cell["coordinate"] == "C1")
        self.assertEqual(formula["value"], "=SUM(B2)")
        self.assertEqual(pages[1]["text"], "")
        self.assertEqual(pages[2]["blocks"][1]["cells"][1]["value"], 3.5)

    def test_default_registry_implements_five_formats(self) -> None:
        registry = create_default_registry()
        expected = {
            ".txt": TxtParser,
            ".md": MdParser,
            ".docx": DocxParser,
            ".pptx": PptxParser,
            ".xlsx": XlsxParser,
        }
        for extension, parser_type in expected.items():
            self.assertIsInstance(registry.find(extension), parser_type)
            self.assertTrue(registry.find(extension).implemented)
            self.assertEqual(registry.status(extension), "implemented")


if __name__ == "__main__":
    unittest.main()